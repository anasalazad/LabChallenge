"""
Tick (or untick) items in the checklist at the top of a weekly worklog,
without opening Word.  (In Word you can also just click the little boxes.)

    python _tools/tick_checklist.py 5              # show lab 5's checklist with item numbers
    python _tools/tick_checklist.py 5 8 9          # tick items 8 and 9 of lab 5
    python _tools/tick_checklist.py 5 --all        # tick everything in lab 5
    python _tools/tick_checklist.py 5 --untick 9   # untick item 9

It flips the checkbox, sets the STATUS column to "Done"/"To do" and updates
the "Progress: x/y items done" line. A backup is saved as <name>.bak.docx.
"""

import os
import re
import shutil
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
DONE_COLOUR = RGBColor(0x37, 0x86, 0x3C)
TODO_COLOUR = RGBColor(0xC0, 0x50, 0x00)


def find_checklist(doc, week):
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip().startswith(f"LAB {week} CHECKLIST"):
            return t
    return None


def is_checked(cell):
    node = cell._tc.find(f".//{{{W14}}}checked")
    return node is not None and node.get(f"{{{W14}}}val") in ("1", "true")


def set_checked(row, value):
    cell0, status = row.cells[0], row.cells[-1]
    node = cell0._tc.find(f".//{{{W14}}}checked")
    node.set(f"{{{W14}}}val", "1" if value else "0")
    for t in cell0._tc.iter(qn("w:t")):
        if t.text in ("☒", "☐"):
            t.text = "☒" if value else "☐"
    runs = status.paragraphs[0].runs
    runs[0].text = "Done" if value else "To do"
    runs[0].font.color.rgb = DONE_COLOUR if value else TODO_COLOUR
    for extra in runs[1:]:
        extra.text = ""


def main(argv):
    if not argv or not argv[0].isdigit():
        print(__doc__)
        return 1
    week = int(argv[0])
    path = os.path.join(REPO_ROOT, f"Lab {week}", f"Lab{week}_Worklog_Week{week}.docx")
    doc = Document(path)
    table = find_checklist(doc, week)
    if table is None:
        print("Couldn't find the checklist table in", path)
        return 1
    rows = table.rows[1:]

    untick = "--untick" in argv
    if "--all" in argv:
        targets = list(range(1, len(rows) + 1))
    else:
        targets = [int(a) for a in argv[1:] if a.isdigit()]

    for n in targets:
        if not 1 <= n <= len(rows):
            print(f"item {n} doesn't exist (there are {len(rows)})")
            return 1
        set_checked(rows[n - 1], not untick)

    done = sum(is_checked(r.cells[0]) for r in rows)
    for p in doc.paragraphs:
        if p.text.startswith("Progress:"):
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = f"Progress: {done}/{len(rows)} items done"
            break

    for i, r in enumerate(rows, 1):
        mark = "[x]" if is_checked(r.cells[0]) else "[ ]"
        text = re.sub(r"\s+", " ", r.cells[1].text.strip())
        print(f"{i:>2}. {mark} {text}")
    print(f"Progress: {done}/{len(rows)}")

    if targets:
        shutil.copyfile(path, path[:-5] + ".bak.docx")
        doc.save(path)
        print("saved", os.path.relpath(path, REPO_ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
