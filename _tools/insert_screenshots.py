"""
Drops your screenshots into the worklogs automatically.

1. Take the screenshot listed in "Lab N/Anas.md" (or in the yellow box in the
   worklog) and save it into that lab's screenshots folder with the name it
   asks for, e.g.   Lab 5/screenshots/5.1_colab_new_notebook.png
   (only the "5.1" prefix really matters - .png, .jpg and .jpeg all work;
   a shot split in two can be saved as 5.1a_xxx.png + 5.1b_xxx.png)
2. From the repo root run:

       python _tools/insert_screenshots.py            # all labs
       python _tools/insert_screenshots.py 5 7        # only labs 5 and 7
       python _tools/insert_screenshots.py --check    # just list what's missing

Every yellow "[SCREENSHOT x.y]" box that has a matching image gets replaced by
the image (the caption under it stays). Boxes with no image yet are left alone,
so you can run this as many times as you like. A backup of each worklog is
written to "<name>.bak.docx" before it is changed (gitignored).
"""

import glob
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp")
MAX_W_IN = 6.2
MAX_H_IN = 8.3
MARKER = re.compile(r"^\[SCREENSHOT\s+(\d+\.\d+[a-z]?)\]")


def find_images(shots_dir, shot_id):
    """All images for one placeholder: '5.3_x.png', '5.3.png', and split shots
    like '5.3a_x.png', '5.3b_x.png' (inserted in a, b, c... order)."""
    if not os.path.isdir(shots_dir):
        return []
    pattern = re.compile(re.escape(shot_id) + r"([a-z])?(?:[ _-].*)?$", re.IGNORECASE)
    found = []
    for name in sorted(os.listdir(shots_dir)):
        if not name.lower().endswith(IMG_EXT):
            continue
        stem = os.path.splitext(name)[0]
        if pattern.match(stem):
            found.append(os.path.join(shots_dir, name))
    return found


def picture_width(path):
    with Image.open(path) as im:
        w, h = im.size
    width = MAX_W_IN
    if h / w * width > MAX_H_IN:
        width = MAX_H_IN * w / h
    return Inches(width)


def process(docx_path, check_only=False):
    lab_dir = os.path.dirname(docx_path)
    shots_dir = os.path.join(lab_dir, "screenshots")
    doc = Document(docx_path)
    inserted, missing = [], []
    for table in list(doc.tables):
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        text = table.rows[0].cells[0].paragraphs[0].text.strip()
        m = MARKER.match(text)
        if not m:
            continue
        shot_id = m.group(1)
        imgs = find_images(shots_dir, shot_id)
        if not imgs:
            missing.append(shot_id)
            continue
        inserted.append((shot_id, ", ".join(os.path.basename(i) for i in imgs)))
        if check_only:
            continue
        for img in imgs:
            p_el = OxmlElement("w:p")
            table._tbl.addprevious(p_el)
            para = Paragraph(p_el, table._parent)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.keep_with_next = True   # keep the caption with the image
            para.add_run().add_picture(img, width=picture_width(img))
        table._tbl.getparent().remove(table._tbl)

    if inserted and not check_only:
        backup = docx_path[:-5] + ".bak.docx"
        shutil.copyfile(docx_path, backup)
        doc.save(docx_path)
    return inserted, missing


def main(argv):
    check_only = "--check" in argv
    labs = [a for a in argv if a.isdigit()]
    pattern = os.path.join(REPO_ROOT, "Lab *", "Lab*_Worklog_Week*.docx")
    files = sorted(f for f in glob.glob(pattern) if not f.endswith(".bak.docx"))
    if labs:
        files = [f for f in files
                 if os.path.basename(os.path.dirname(f)).split()[-1] in labs]
    if not files:
        print("No worklogs found - run this from the repo (python _tools/insert_screenshots.py)")
        return 1
    for f in files:
        ins, miss = process(f, check_only)
        rel = os.path.relpath(f, REPO_ROOT)
        verb = "ready to insert" if check_only else "inserted"
        print(f"\n{rel}")
        print(f"  {verb}: " + (", ".join(f"{i} ({n})" for i, n in ins) or "-"))
        print("  still missing: " + (", ".join(miss) or "nothing - all screenshots in!"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
