"""
Builds my Lab 8 deliverable: the ONE-PAGE summary (results table, gradient plot,
attention plot, answers to the 3 checkpoint questions).

    python make_summary.py            (run inside the "Lab 8" folder)

It reads results.json + figures/ written by Lab8_RNN_LSTM_Attention.ipynb, so if
the notebook is re-run, just run this again. Output: Lab8_Summary.docx
(then File > Save As > PDF in Word, or: soffice --headless --convert-to pdf Lab8_Summary.docx)
"""
import json
import os

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
STUDENT = "Anas Al Azad"
STUDENT_ID = ""          # <- your student ID

# ---- my answers to the checkpoint questions (2-3 sentences each) -------------
ANSWERS = [
    ("1. In Part 1, which input step has the smallest gradient? What does this mean for learning "
     "the first digit?",
     "{Q1}"),
    ("2. In Part 2, runs with the same settings give different accuracy. What does this tell you "
     "about training RNNs on long sequences?",
     "{Q2}"),
    ("3. In Part 3, accuracy drops when you remove positional encoding. Why?",
     "{Q3}"),
]

ROWS = [("rnn", "RNN"), ("lstm", "LSTM"), ("gru", "GRU"),
        ("rnn-128", "RNN, 128 units"), ("lstm-128", "LSTM, 128 units"), ("gru-128", "GRU, 128 units"),
        ("attn", "Self-attention"), ("attn_nopos", "Attention, no pos. enc.")]
TS = [5, 10, 20, 30, 60]


def load():
    with open(os.path.join(HERE, "results.json")) as f:
        data = json.load(f)
    res = {}
    for key, accs in data["results"].items():
        name, T = key.split("|")
        res[(name, int(T))] = accs
    return data, res


def sci(x):
    """1.5e+13 -> '1.5×10¹³' (looks nicer in the summary)"""
    mant, exp = f"{x:.1e}".split("e")
    sup = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")
    return f"{mant}×10{str(int(exp)).translate(sup)}"


def order_blind_best(T, n=200_000, seed=0):
    """Best accuracy possible without knowing the order: guess the most common digit."""
    X = np.random.default_rng(seed).integers(0, 10, size=(n, T))
    counts = np.stack([(X == d).sum(1) for d in range(10)], axis=1)
    return (counts.max(1) / T).mean()


def fmt_answers(data, res):
    """The answers quote my actual numbers, so they're filled in from results.json."""
    g = data["grad_ratio"]
    m = lambda k, T: np.mean(res[(k, T)])
    runs = lambda k, T: " / ".join(f"{a:.2f}" for a in res[(k, T)])
    w0 = data["attention_weight_pos0"]
    vals = dict(
        Q1=(f"Step 0 – the very first input – has the smallest gradient for all three layers: the "
            f"gradient at the last step is about {sci(g['RNN'])} times bigger for the RNN, {sci(g['LSTM'])} for "
            f"the LSTM and {sci(g['GRU'])} for the GRU. But the first digit is exactly what the label is, so the "
            f"weights get almost no training signal about how to store it. That's why the plain RNN drops to "
            f"chance ({m('rnn', 20):.2f}) from T = 20, while the gated LSTM/GRU hold on a bit longer "
            f"(fine at T = 10)."),
        Q2=(f"Only the seed changes (initial weights + data order), yet the results jump around a lot: the RNN "
            f"at T = 10 got {runs('rnn', 10)} and the LSTM at T = 20 got {runs('lstm', 20)}. Because the "
            f"gradient from the early steps is so tiny, whether a run ever finds the link between step 1 and "
            f"the label is basically luck. So training RNNs on long sequences is unstable – you need "
            f"several seeds and should look at the spread, not just one number."),
        Q3=(f"Self-attention on its own doesn't know about order – it compares steps by content only, so "
            f"without positional encoding the sequence is just a bag of digits and the model can't tell which "
            f"one came first. The best it can do is guess the most common digit, which I calculated gives "
            f"≈{order_blind_best(10):.2f} for T = 10 and ≈{order_blind_best(30):.2f} for T = 30 – "
            f"basically the {m('attn_nopos', 10):.2f} and {m('attn_nopos', 30):.2f} I got. With positional "
            f"encoding every step gets its own sin/cos tag, so the last position looks straight at position 0 "
            f"(weights {w0[0]:.2f} / {w0[1]:.2f}) and gets {m('attn', 10):.2f}."),
    )
    return [(q, a.format(**vals)) for q, a in ANSWERS]


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def insert_ordered(parent, child, order):
    """Word wants child elements in schema order - insert in the right spot."""
    name = child.tag.split("}")[1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    later = order[order.index(name) + 1:]
    for i, sib in enumerate(parent):
        if sib.tag.split("}")[1] in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def fixed_layout(table):
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    insert_ordered(table._tbl.tblPr, lay, TBLPR_ORDER)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    insert_ordered(tcpr, shd, TCPR_ORDER)


def run(p, text, size=9, bold=False, italic=False, colour=None):
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.name = "Arial"
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    if colour:
        r.font.color.rgb = colour
    return r


def para(doc, before=0, after=2, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    return p


def build():
    data, res = load()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    sec.left_margin = sec.right_margin = Cm(1.6)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)

    p = para(doc, after=0)
    run(p, "COS30018 Intelligent Systems – Lab 8: RNN, LSTM and Attention", 13, bold=True)
    p = para(doc, after=6)
    run(p, f"One-page summary · {STUDENT} · Student ID: ", 9, colour=RGBColor(0x40, 0x40, 0x40))
    run(p, STUDENT_ID or "________", 9, bold=bool(STUDENT_ID))
    run(p, " · Notebook: Lab8_RNN_LSTM_Attention.ipynb", 9, colour=RGBColor(0x40, 0x40, 0x40))

    # ---- results table ---------------------------------------------------------
    p = para(doc, before=2, after=2)
    run(p, "Results (Parts 2 and 3)", 10.5, bold=True)
    run(p, "  – mean test accuracy over 3 seeds; task = recall the first of T random digits; "
           "chance = 0.10", 8.5, italic=True)
    table = doc.add_table(rows=1 + len(ROWS), cols=1 + len(TS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for c, text in zip(hdr, ["Model"] + [f"T = {T}" for T in TS]):
        shade(c, "DEEAF6")
        cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(cp, text, 9, bold=True)
    for r, (key, label) in zip(table.rows[1:], ROWS):
        cp = r.cells[0].paragraphs[0]
        run(cp, label, 9, bold=key in ("attn",))
        for c, T in zip(r.cells[1:], TS):
            cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if (key, T) in res:
                v = np.mean(res[(key, T)])
                colour = (RGBColor(0x1E, 0x7B, 0x34) if v >= 0.9 else
                          RGBColor(0xB0, 0x1E, 0x1E) if v <= 0.2 else None)
                run(cp, f"{v:.2f}", 9, bold=v >= 0.9, colour=colour)
            else:
                run(cp, "–", 8.5, colour=RGBColor(0x99, 0x99, 0x99))
    table.autofit = False
    fixed_layout(table)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), [5.0] + [2.2] * len(TS)):
        gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
    for row in table.rows:
        row.cells[0].width = Cm(5.0)
        for c in row.cells[1:]:
            c.width = Cm(2.2)
        for c in row.cells:
            for cp in c.paragraphs:
                cp.paragraph_format.space_before = Pt(0.5)
                cp.paragraph_format.space_after = Pt(0.5)
    p = para(doc, before=2, after=4)
    run(p, "Experiment (32 → 128 hidden units, T = 20) and the ablation (no positional encoding) are "
           "the extra rows. Green = ≥ 0.90, red = ≤ 0.20 (about chance).", 7.5, italic=True,
        colour=RGBColor(0x55, 0x55, 0x55))

    # ---- plots -------------------------------------------------------------------
    p = para(doc, before=2, after=2)
    run(p, "Gradient plot (Part 1) and attention plot (Part 3)", 10.5, bold=True)
    pt = doc.add_table(rows=1, cols=2)
    pt.alignment = WD_TABLE_ALIGNMENT.CENTER
    pt.autofit = False
    fixed_layout(pt)
    for c, (img, w, cw) in zip(pt.rows[0].cells, [("figures/8_1_gradient_plot.png", 2.68, 2.85),
                                                   ("figures/8_2_attention_plot.png", 3.72, 3.95)]):
        c.width = Inches(cw)
        cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run().add_picture(os.path.join(HERE, img), width=Inches(w))
    for gc, cw in zip(pt._tbl.tblGrid.findall(qn("w:gridCol")), (2.85, 3.95)):
        gc.set(qn("w:w"), str(int(cw * 1440)))
    p = para(doc, before=1, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    g = data["grad_ratio"]
    w0 = data["attention_weight_pos0"]
    run(p, f"Left: gradient size at each of 50 input steps (log scale) – last/first ratio "
           f"RNN {sci(g['RNN'])}, LSTM {sci(g['LSTM'])}, GRU {sci(g['GRU'])}.  Right: attention of the last "
           f"position at T = 20, averaged over 200 sequences – weight on position 0: head 0 = {w0[0]:.2f}, "
           f"head 1 = {w0[1]:.2f}.", 7.5, italic=True, colour=RGBColor(0x55, 0x55, 0x55))

    # ---- answers -------------------------------------------------------------------
    p = para(doc, before=2, after=2)
    run(p, "Checkpoint questions", 10.5, bold=True)
    for q, a in fmt_answers(data, res):
        p = para(doc, before=2, after=0)
        run(p, q, 10, bold=True)
        p = para(doc, before=0, after=4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run(p, a, 10)

    out = os.path.join(HERE, "Lab8_Summary.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build()
